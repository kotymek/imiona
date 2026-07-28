"""Pobiera najnowsze listy pierwszych imion z dane.gov.pl i tworzy plik dla strony."""
from __future__ import annotations

import io
import json
import re
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

API_URL = "https://api.dane.gov.pl/1.4/datasets/1667/resources?per_page=100"
ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "data" / "names.json"


def request(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "kotymek.github.io/1.0"})
    with urllib.request.urlopen(req, timeout=60) as response:
        return response.read()


def latest_resources() -> tuple[str, list[dict]]:
    payload = json.loads(request(API_URL))
    resources = [
        item["attributes"] for item in payload["data"]
        if item["attributes"].get("format", "").lower() == "xlsx"
        and "imię pierwsze" in item["attributes"].get("title", "").lower()
        and ("męskich" in item["attributes"]["title"].lower() or "żeńskich" in item["attributes"]["title"].lower())
    ]
    latest_date = max(item["data_date"] for item in resources)
    latest = [item for item in resources if item["data_date"] == latest_date]
    if len(latest) != 2:
        raise RuntimeError(f"Oczekiwano 2 najnowszych zasobów, znaleziono: {len(latest)}")
    return latest_date, latest


def rows(resource: dict) -> list[dict]:
    workbook = load_workbook(io.BytesIO(request(resource["link"])), read_only=True, data_only=True)
    sheet = workbook.active
    gender = "M" if "męskich" in resource["title"].lower() else "K"
    result = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name, _, count = row[:3]
        if not name or count is None:
            continue
        clean_name = re.sub(r"\s+", " ", repair_encoding(str(name)).strip().upper())
        result.append({"name": clean_name, "gender": gender, "count": int(count)})
    return result


def repair_encoding(value: str) -> str:
    """Portalowe XLSX-y mają tekst UTF-8 omyłkowo odczytany jako Windows-1250."""
    try:
        return value.encode("cp1250").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return value


def main() -> None:
    data_date, resources = latest_resources()
    names = [entry for resource in resources for entry in rows(resource)]
    names.sort(key=lambda item: (-item["count"], item["name"], item["gender"]))
    payload = {
        "date": data_date,
        "updated": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source": "https://dane.gov.pl/pl/dataset/1667",
        "names": names,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Zapisano {len(names)} pozycji z dnia {data_date}.")


if __name__ == "__main__":
    main()
